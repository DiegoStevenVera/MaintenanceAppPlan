"""Import the warehouse tool inventory workbook and its label photographs."""

import hashlib
import mimetypes
import shutil
import unicodedata
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import UUID, uuid4, uuid5

from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from modules.asset_management.infrastructure.postgres.domain_models import InventoryLocationRecord
from modules.maintenance_execution.infrastructure.postgres.tool_inventory_item_models import (
    ToolInventoryItemImageRecord,
    ToolInventoryItemRecord,
)
from modules.maintenance_execution.infrastructure.postgres.tool_inventory_models import (
    ToolInventoryRecord,
)
from modules.maintenance_execution.infrastructure.postgres.tool_models import (
    ToolCatalogItemRecord,
    ToolRecord,
)
from shared_kernel.storage import storage_key


SOURCE_NAME = "INVENTARIO_HERRAMIENTAS.xlsx"
TOOL_IMAGE_DIRECTORY = "tool-inventory"
INVENTORY_NAMESPACE = UUID("0f1d9b06-c1b8-4e1f-bf55-e4f61cb8b4a0")
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".heic"}


@dataclass
class ToolInventoryImportSummary:
    rows_read: int = 0
    rows_imported: int = 0
    rows_skipped: int = 0
    images_imported: int = 0
    rows_without_images: int = 0
    catalog_items_retired: int = 0
    location_name: str | None = None


def _normalized(value: object) -> str:
    text = str(value or "").strip()
    text = unicodedata.normalize("NFKD", text)
    return "".join(char for char in text if not unicodedata.combining(char)).upper()


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _decimal(value: object, *, row_number: int) -> Decimal:
    text = _text(value) or "0"
    try:
        result = Decimal(text.replace(",", "."))
    except InvalidOperation as error:
        raise ValueError(f"Fila {row_number}: CANTIDAD no es numérica: {text}") from error
    if result < 0:
        raise ValueError(f"Fila {row_number}: CANTIDAD no puede ser negativa")
    return result


def _classification_code(classification: str | None) -> str:
    classification_key = _normalized(classification)
    if "CONSUMIBLE" in classification_key:
        return "CONSUMABLE"
    if "LLAVE" in classification_key or "ACCESO" in classification_key:
        return "ACCESS_KEY"
    if "EQUIPO" in classification_key:
        return "EQUIPMENT"
    return "MANUAL_TOOL"


def _photo_directory(photos_root: Path, label: str | None) -> Path | None:
    if not label:
        return None
    candidates = [photos_root / label]
    if label.isdigit():
        candidates.append(photos_root / label.zfill(3))
        stripped = label.lstrip("0")
        if stripped:
            candidates.append(photos_root / stripped)
    return next((candidate for candidate in candidates if candidate.is_dir()), None)


async def _inventory_location(
    session: AsyncSession,
    value: str | None,
) -> InventoryLocationRecord | None:
    if not value:
        return None
    locations = (await session.scalars(select(InventoryLocationRecord))).all()
    target = _normalized(value)
    aliases = {
        "ALMACEN HITACHI": "ALMACENAMIENTO MANTTO HITACHI",
        "ALMACENAMIENTO HITACHI": "ALMACENAMIENTO MANTTO HITACHI",
        "ALMACEN SPV": "ALMACENAMIENTO SPV",
    }
    target = aliases.get(target, target)
    for location in locations:
        if _normalized(location.name) == target:
            return location
    available = ", ".join(sorted(location.name for location in locations))
    raise ValueError(
        f"No existe una ubicación de almacenamiento para '{value}'. "
        f"Ubicaciones disponibles: {available}"
    )


async def _catalog_item(
    session: AsyncSession,
    *,
    name: str,
    classification: str | None,
) -> ToolCatalogItemRecord:
    code = _classification_code(classification)
    item = await session.scalar(
        select(ToolCatalogItemRecord).where(
            func.lower(ToolCatalogItemRecord.name) == name.lower(),
            ToolCatalogItemRecord.category == code,
        )
    )
    if item:
        if not item.is_active:
            item.is_active = True
        return item
    item = ToolCatalogItemRecord(
        id=uuid4(),
        name=name,
        category=code,
        default_unit="unidad",
        requires_identified_unit=code == "EQUIPMENT",
        is_active=True,
    )
    session.add(item)
    await session.flush()
    return item


async def import_tool_inventory(
    session: AsyncSession,
    workbook_path: Path,
    photos_root: Path,
    *,
    replace: bool = False,
    dry_run: bool = False,
) -> ToolInventoryImportSummary:
    from openpyxl import load_workbook

    workbook_path = workbook_path.resolve()
    photos_root = photos_root.resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError(workbook_path)
    if not photos_root.is_dir():
        raise FileNotFoundError(photos_root)

    worksheet = load_workbook(workbook_path, read_only=True, data_only=True).active
    header_values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = {_normalized(value): index for index, value in enumerate(header_values) if value}
    required = {
        "CODIGO_MODELO",
        "GABINETE",
        "ETIQUETA",
        "CANTIDAD",
        "NOMBRE",
        "DESCRIPCION",
        "ARMARIO",
        "ESTADO",
        "EMPRESA",
        "CLASIFICACION",
        "CATEGORIA",
        "UBICACION",
        "CODIGO",
        "OBSERVACIONES",
    }
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"Faltan columnas en el Excel: {', '.join(missing)}")

    summary = ToolInventoryImportSummary()
    imported_catalog_ids: set[UUID] = set()
    image_root = settings.resolved_attachment_root / TOOL_IMAGE_DIRECTORY
    if replace:
        await session.execute(
            delete(ToolInventoryItemRecord).where(
                ToolInventoryItemRecord.source_file == SOURCE_NAME
            )
        )
        if not dry_run and image_root.exists():
            shutil.rmtree(image_root)

    if not dry_run:
        image_root.mkdir(parents=True, exist_ok=True)
    for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), 2):
        if not any(value not in (None, "") for value in values):
            summary.rows_skipped += 1
            continue
        summary.rows_read += 1

        def value(column: str) -> object:
            return values[headers[column]] if headers[column] < len(values) else None

        name = _text(value("NOMBRE"))
        if not name:
            raise ValueError(f"Fila {row_number}: NOMBRE es obligatorio")
        location_text = _text(value("UBICACION"))
        location = await _inventory_location(session, location_text)
        catalog = await _catalog_item(
            session,
            name=name,
            classification=_text(value("CLASIFICACION")),
        )
        imported_catalog_ids.add(catalog.id)
        item_id = uuid5(INVENTORY_NAMESPACE, f"{SOURCE_NAME}:{row_number}")
        item = await session.get(ToolInventoryItemRecord, item_id)
        if item is None:
            item = ToolInventoryItemRecord(id=item_id, source_file=SOURCE_NAME, source_row_number=row_number)
        item.catalog_item_id = catalog.id
        item.model_code = _text(value("CODIGO_MODELO"))
        item.tool_code = _text(value("CODIGO"))
        item.label = _text(value("ETIQUETA"))
        item.quantity = _decimal(value("CANTIDAD"), row_number=row_number)
        item.name = name
        item.description = _text(value("DESCRIPCION"))
        item.cabinet = _text(value("ARMARIO"))
        item.cabinet_detail = _text(value("GABINETE"))
        item.company = _text(value("EMPRESA"))
        item.classification = _text(value("CLASIFICACION"))
        item.category = _text(value("CATEGORIA"))
        item.inventory_location_id = location.id if location else None
        item.location_text = location_text
        item.status = _text(value("ESTADO"))
        item.observations = _text(value("OBSERVACIONES"))
        if catalog.requires_identified_unit and item.tool_code:
            tool = await session.scalar(
                select(ToolRecord).where(ToolRecord.serial_number == item.tool_code)
            )
            if tool is None:
                tool = ToolRecord(
                    id=uuid4(),
                    serial_number=item.tool_code,
                    name=name,
                    availability_status="AVAILABLE",
                )
            tool.catalog_item_id = catalog.id
            tool.name = name
            tool.model = item.model_code
            tool.brand = item.company
            tool.tool_type = item.category or "INSTRUMENTO"
            tool.current_location = location.name if location else location_text
            tool.availability_status = (
                "AVAILABLE"
                if _normalized(item.status) == "DISPONIBLE" and item.quantity > 0
                else "UNAVAILABLE"
            )
            tool.is_active = True
            session.add(tool)
            await session.flush()
            item.tool_id = tool.id
            item.quantity = Decimal(1)
            stock = await session.scalar(
                select(ToolInventoryRecord).where(ToolInventoryRecord.tool_id == tool.id)
            )
            if stock is None:
                is_available = tool.availability_status == "AVAILABLE"
                session.add(
                    ToolInventoryRecord(
                        id=uuid4(),
                        catalog_item_id=catalog.id,
                        tool_id=tool.id,
                        location=tool.current_location or "Sin ubicación",
                        available=1 if is_available else 0,
                        unavailable=0 if is_available else 1,
                    )
                )
        session.add(item)
        await session.flush()
        await session.execute(
            delete(ToolInventoryItemImageRecord).where(
                ToolInventoryItemImageRecord.inventory_item_id == item.id
            )
        )

        photos = sorted(
            photo
            for photo in (_photo_directory(photos_root, item.label) or Path()).iterdir()
            if photo.is_file() and photo.suffix.lower() in IMAGE_EXTENSIONS
        ) if _photo_directory(photos_root, item.label) else []
        destination = image_root / str(item.id)
        for photo in photos:
            target = destination / photo.name
            content = photo.read_bytes()
            if not dry_run:
                destination.mkdir(parents=True, exist_ok=True)
                shutil.copy2(photo, target)
            session.add(
                ToolInventoryItemImageRecord(
                    id=uuid4(),
                    inventory_item_id=item.id,
                    file_name=photo.name,
                    file_reference=storage_key(target, settings.resolved_attachment_root),
                    media_type=mimetypes.guess_type(photo.name)[0] or "application/octet-stream",
                    byte_size=len(content),
                    checksum=hashlib.sha256(content).hexdigest(),
                )
            )
            summary.images_imported += 1
        if not photos:
            summary.rows_without_images += 1
        summary.rows_imported += 1
        summary.location_name = location.name if location else summary.location_name

    if replace:
        # Retire only generic manual/consumable definitions absent from the new
        # warehouse source. Existing report history and checklist snapshots keep
        # their foreign keys, but the obsolete definitions disappear from active
        # administration screens.
        old_catalog = (
            await session.scalars(
                select(ToolCatalogItemRecord).where(
                    ToolCatalogItemRecord.category.in_(("MANUAL_TOOL", "CONSUMABLE")),
                    ToolCatalogItemRecord.is_active.is_(True),
                )
            )
        ).all()
        for catalog in old_catalog:
            if catalog.id not in imported_catalog_ids:
                catalog.is_active = False
                summary.catalog_items_retired += 1

    await session.flush()
    return summary
